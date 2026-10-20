"""Generate synthetic test workbooks for input-agnostic solver tests.

Run once:  python tests/fixtures/make_synthetic_workbooks.py
"""
from pathlib import Path
import openpyxl

OUT = Path(__file__).parent


def make_ttest_nonadjacent():
    """Two-sample t-test data in non-adjacent columns C and E (not A/B)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TTest"
    ws["A1"] = "Two-sample t-test: compare group A (col C) vs group B (col E)"
    ws["C1"] = "Group A"
    ws["E1"] = "Group B"
    group_a = [5.1, 4.8, 5.3, 5.0, 4.9, 5.2, 4.7, 5.1]
    group_b = [4.3, 4.0, 4.5, 4.1, 4.4, 4.2, 4.6, 4.0]
    for i, (a, b) in enumerate(zip(group_a, group_b), start=2):
        ws[f"C{i}"] = a
        ws[f"E{i}"] = b
    wb.save(OUT / "ttest_nonadjacent.xlsx")
    print("Created ttest_nonadjacent.xlsx")


def make_descriptive_col_b():
    """Descriptive statistics data in column B (not A)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stats"
    ws["A1"] = "Descriptive statistics exercise"
    ws["B1"] = "Values"
    data = [12.5, 14.3, 11.8, 13.2, 15.0, 12.9, 14.1, 13.7, 12.1, 14.8]
    for i, v in enumerate(data, start=2):
        ws[f"B{i}"] = v
    wb.save(OUT / "descriptive_col_b.xlsx")
    print("Created descriptive_col_b.xlsx")


def make_onesample_with_mu0():
    """One-sample hypothesis test with mu0 specified in cell text."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "HypTest"
    ws["A1"] = "One-sample t-test: H0: mean = 50.0"
    ws["D1"] = "Sample"
    data = [48.2, 51.5, 49.8, 52.1, 50.3, 47.9, 51.0, 49.5, 50.8, 48.7]
    for i, v in enumerate(data, start=2):
        ws[f"D{i}"] = v
    wb.save(OUT / "onesample_mu0.xlsx")
    print("Created onesample_mu0.xlsx")


if __name__ == "__main__":
    make_ttest_nonadjacent()
    make_descriptive_col_b()
    make_onesample_with_mu0()
